import docx
import pytest
from openpyxl import Workbook


from codelens.index import build_index
from codelens.reader import read_file
from codelens.search import search_all


def test_word_file_text_and_tables(tmp_path):
    path = tmp_path / "invoice.docx"
    document = docx.Document()
    document.add_paragraph("Invoice for March")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Total"
    table.rows[0].cells[1].text = "4500"
    document.save(str(path))

    text = read_file(path)
    assert "Invoice for March" in text
    assert "Total" in text and "4500" in text


def test_excel_file_cells(tmp_path):
    path = tmp_path / "sales.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["region", "amount"])
    sheet.append(["north", 120])
    book.save(str(path))


    text = read_file(path)
    assert "region amount" in text
    assert "north 120" in text


def test_pdf_file_text(tmp_path):
    canvas_module = pytest.importorskip("reportlab.pdfgen.canvas")
    path = tmp_path / "notes.pdf"
    pdf = canvas_module.Canvas(str(path))
    pdf.drawString(72, 720, "Quarterly report on search engines")
    pdf.save()

    assert "Quarterly report" in read_file(path)


def test_damaged_documents_give_none(tmp_path):
    for name in ("broken.pdf", "broken.docx", "broken.xlsx"):
        (tmp_path / name).write_bytes(b"this is not a real document")
        assert read_file(tmp_path / name) is None



def test_plain_text_still_works(tmp_path):
    path = tmp_path / "a.py"
    path.write_text("def run(): pass", encoding="utf-8")
    assert read_file(path) == "def run(): pass"



def test_text_is_cut_at_the_safety_limit(tmp_path, monkeypatch):
    path = tmp_path / "big.txt"
    path.write_text("word " * 100, encoding="utf-8")
    monkeypatch.setattr("codelens.config.MAX_TEXT_CHARS", 20)
    assert len(read_file(path)) == 20


    
def test_documents_are_found_by_a_search(tmp_path):
    document = docx.Document()
    document.add_paragraph("Invoice for March")
    document.save(str(tmp_path / "invoice.docx"))
    (tmp_path / "notes.txt").write_text("nothing here", encoding="utf-8")


    index = build_index(str(tmp_path))
    hits = search_all(index, "invoice march")
    assert [index.docs[d]["path"] for d in hits] == ["invoice.docx"]