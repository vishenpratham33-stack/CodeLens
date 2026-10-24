import base64
from pathlib import Path

import docx
import pytest
from fastapi.testclient import TestClient

from codelens.api.server import clean_upload_name, create_app
from codelens.library import Library

SECRET = b"test-secret-test-secret-test-secret"


def sign_up(client, name="alice", password="a long password"):
    client.post("/api/register", json={"username": name, "password": password})
    reply = client.post("/api/login", json={"username": name, "password": password})
    return {"Authorization": "Bearer " + reply.json()["token"]}


@pytest.fixture
def docs(tmp_path):
    folder = tmp_path / "Documents"
    folder.mkdir()
    (folder / "cv.txt").write_text("curriculum vitae of a python developer", encoding="utf-8")
    word = docx.Document()
    word.add_paragraph("Electricity bill for March, amount due 4500")
    word.save(str(folder / "bill.docx"))
    return folder


@pytest.fixture
def opened():
    return []


@pytest.fixture
def make_client(tmp_path, docs, opened):
    def build(**client_options):
        library = Library([docs], tmp_path / "store")
        library.sync()
        app = create_app(tmp_path / "data", secret=SECRET, library=library,
                         file_opener=opened.append)
        return TestClient(app, **client_options), library
    return build


def test_me_says_which_mode(make_client):
    client, _ = make_client()
    headers = sign_up(client)
    info = client.get("/api/me", headers=headers).json()
    assert info["mode"] == "library"
    assert info["can_open"] is False                  # the test client is not "this computer"


def test_files_are_found_without_any_upload(make_client):
    client, _ = make_client()
    headers = sign_up(client)
    reply = client.get("/api/search?q=electricity bill", headers=headers).json()
    assert reply["total"] == 1
    assert reply["results"][0]["path"].endswith("/bill.docx")
    assert "Electricity bill" in reply["results"][0]["snippet"]


def test_typo_and_suggestions_work_in_my_files_mode(make_client):
    client, _ = make_client()
    headers = sign_up(client)
    typo = client.get("/api/search?q=developr", headers=headers).json()
    assert typo["corrections"] == [{"typed": "developr", "used": "developer"}]
    words = client.get("/api/suggest?prefix=curr", headers=headers).json()["suggestions"]
    assert words[0]["word"] == "curriculum"


def test_library_status(make_client):
    client, _ = make_client()
    headers = sign_up(client)
    status = client.get("/api/library", headers=headers).json()
    assert status["state"] == "ready" and status["files"] == 2
    assert status["roots"][0].endswith("Documents")


def test_uploads_and_deletes_are_switched_off(make_client):
    client, _ = make_client()
    headers = sign_up(client)
    assert client.post("/api/files", json={"name": "x.txt", "content": "no"}, headers=headers).status_code == 403
    assert client.delete("/api/files/cv.txt", headers=headers).status_code == 403
    assert client.get("/api/files", headers=headers).json() == {"files": []}


def test_open_works_only_from_this_computer(make_client, opened):
    remote, library = make_client()
    headers = sign_up(remote)
    path = next(m["path"] for m in library.index.docs.values() if m["path"].endswith("cv.txt"))
    assert remote.post("/api/open", json={"path": path}, headers=headers).status_code == 403
    assert opened == []

    local, _ = make_client(client=("127.0.0.1", 50000), base_url="http://127.0.0.1:8000")
    headers = sign_up(local, "bob", "another long one")
    assert local.get("/api/me", headers=headers).json()["can_open"] is True
    assert local.post("/api/open", json={"path": path}, headers=headers).status_code == 200
    assert [Path(p) for p in opened] == [Path(path)]


def test_open_refuses_a_tunnel_even_from_127_0_0_1(make_client, opened):
    client, library = make_client(client=("127.0.0.1", 50000), base_url="http://127.0.0.1:8000")
    headers = sign_up(client)
    path = next(iter(library.index.docs.values()))["path"]
    tunnelled = dict(headers, **{"X-Forwarded-For": "203.0.113.9"})
    assert client.post("/api/open", json={"path": path}, headers=tunnelled).status_code == 403
    other_host = dict(headers, **{"Host": "laptop.example.ts.net"})
    assert client.post("/api/open", json={"path": path}, headers=other_host).status_code == 403
    assert opened == []


def test_only_indexed_files_can_be_opened_or_downloaded(make_client, tmp_path, opened):
    client, library = make_client(client=("127.0.0.1", 50000), base_url="http://127.0.0.1:8000")
    headers = sign_up(client)
    secret = tmp_path / "secret.txt"
    secret.write_text("not indexed", encoding="utf-8")
    for bad in [str(secret), "/etc/passwd", "C:/Windows/win.ini", "../secret.txt"]:
        assert client.post("/api/open", json={"path": bad}, headers=headers).status_code == 404
        assert client.get("/api/download", params={"path": bad}, headers=headers).status_code == 404
    assert opened == []


def test_download_an_indexed_file(make_client):
    client, library = make_client()
    headers = sign_up(client)
    path = next(m["path"] for m in library.index.docs.values() if m["path"].endswith("cv.txt"))
    reply = client.get("/api/download", params={"path": path}, headers=headers)
    assert reply.status_code == 200
    assert b"curriculum vitae" in reply.content
    assert "cv.txt" in reply.headers["content-disposition"]
    assert client.get("/api/download", params={"path": path}).status_code == 401


# ---------- the upload fix (hosted mode) ----------
@pytest.fixture
def hosted(tmp_path):
    return TestClient(create_app(tmp_path / "data", secret=SECRET))


def test_names_with_spaces_and_brackets_are_cleaned(hosted):
    headers = sign_up(hosted)
    reply = hosted.post("/api/files", json={"name": "My CV (final).txt", "content": "hello"}, headers=headers)
    assert reply.status_code == 201
    assert reply.json()["saved"] == "My_CV_final_.txt"


def test_a_word_file_can_be_uploaded_as_base64(hosted):
    headers = sign_up(hosted)
    word = docx.Document()
    word.add_paragraph("Quarterly invoice summary")
    import io
    buffer = io.BytesIO()
    word.save(buffer)
    body = {"name": "report.docx", "content_base64": base64.b64encode(buffer.getvalue()).decode()}
    assert hosted.post("/api/files", json=body, headers=headers).status_code == 201
    assert hosted.get("/api/search?q=invoice", headers=headers).json()["total"] == 1


def test_bad_base64_is_refused(hosted):
    headers = sign_up(hosted)
    reply = hosted.post("/api/files", json={"name": "a.pdf", "content_base64": "###"}, headers=headers)
    assert reply.status_code == 400


def test_name_cleaning_rules():
    assert clean_upload_name("a b.txt") == "a_b.txt"
    for bad in ["../x.txt", "a/b.txt", "a\\b.txt", ".hidden.txt", "C:evil.txt", "....", "***"]:
        assert clean_upload_name(bad) is None


def test_service_worker_and_page_are_served(tmp_path):
    web = tmp_path / "web"
    web.mkdir()
    (web / "index.html").write_text("<html>hi</html>", encoding="utf-8")
    (web / "sw.js").write_text("// worker", encoding="utf-8")
    client = TestClient(create_app(tmp_path / "data", secret=SECRET, web_dir=web))
    assert client.get("/").status_code == 200
    reply = client.get("/sw.js")
    assert reply.status_code == 200 and "javascript" in reply.headers["content-type"]
