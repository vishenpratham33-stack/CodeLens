from pathlib import Path
from typing import Optional

from openpyxl import load_workbook


def extract_xlsx(path: Path) -> Optional[str]:
    try:
        book = load_workbook(str(path), read_only=True, data_only=True)
        lines = []
        for sheet in book.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [str(value) for value in row if value is not None]
                if cells:
                    lines.append(" ".join(cells))
        book.close()
    except Exception:
        return None
    return "\n".join(lines)