from pathlib import Path
from typing import Optional

import docx


def extract_docx(path: Path) -> Optional[str]:
    try:
        document = docx.Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
    except Exception:
        return None
    return "\n".join(parts)